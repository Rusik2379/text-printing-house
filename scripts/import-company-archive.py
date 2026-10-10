"""Import company text from an extracted TEKST_HTML_site archive (requires lxml)."""
from pathlib import Path
from datetime import date
from lxml import html, etree
from copy import deepcopy
import argparse, json, re, hashlib

ROOT=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('source',type=Path,help='Extracted TEKST_HTML_site directory')
parser.add_argument('--archive',type=Path,help='Original ZIP/RAR for the provenance hash')
parser.add_argument('--seo',type=Path,help='Latest customer SEO workbook (read only)')
parser.add_argument('--imported',default=date.today().isoformat(),help='Content import date (YYYY-MM-DD)')
args=parser.parse_args()
source=args.source.resolve()
catalog_text=(ROOT/'catalog.js').read_text(encoding='utf-8')
catalog=json.loads(catalog_text[catalog_text.index('['):catalog_text.rindex(']')+1])
services={r['url']:r for r in catalog}
legacy=json.loads((ROOT/'company-articles.json').read_text(encoding='utf-8'))['articles']
previous=json.loads((ROOT/'company-content.json').read_text(encoding='utf-8'))
previous_articles={a['source_route']:a for a in previous['articles']}
def norm(text):return re.sub(r'\W+','',text.lower().replace('ё','е'))
legacy_paths={norm(a['title']):a['path'] for a in legacy}
def text(node):return ' '.join(''.join(node.itertext()).split())
def load(path):return html.fromstring((source/path).read_bytes())

ALLOWED={'p','h2','h3','h4','ul','ol','li','strong','b','em','i','br','a','code','table','thead','tbody','tr','th','td','aside','blockquote'}
def clean(node):
    node=deepcopy(node)
    for element in [node,*list(node.iterdescendants())]:
        if element.tag in ('script','style','iframe'):
            element.drop_tree();continue
        if element.tag not in ALLOWED:
            if element is not node:element.drop_tag()
            continue
        route=element.get('data-route') or element.get('href')
        element.attrib.clear()
        if element.tag=='a' and route:
            if route.startswith('/') and not route.startswith('//'):element.set('href','SOURCE:'+route)
            elif route.startswith(('https://','mailto:','tel:')):element.set('href',route)
            else:raise ValueError('Unmapped source link: '+route)
    node.tail=None
    return etree.tostring(node,method='html',encoding='unicode')

pages={}
for p in sorted((source/'company').rglob('index.html')):
    doc=load(p.relative_to(source));main=doc.xpath('//main')[0]
    pages[p.relative_to(source).as_posix()]={
        'title':text(main.xpath('.//h1')[0]),
        'seo_title':doc.xpath('string(//title)'),
        'description':doc.xpath('string(//meta[@name="description"]/@content)'),
        'sections':[{'tag':e.tag,'text':text(e),'html':clean(e)} for e in main.xpath('.//h2|.//h3|.//p') if not e.xpath('ancestor::details')],
        'links':[{'label':text(a),'url':a.get('href')} for a in main.xpath('.//a[starts-with(@href,"https://")]')],
        'table_rows':[[text(c) for c in row.xpath('./th|./td')] for row in main.xpath('.//tr')]
    }
    if p.relative_to(source).as_posix()=='company/delivery/index.html':
        capacity=[]
        for link in main.xpath('.//div[contains(@class,"service-index")]/a'):
            match=re.fullmatch(r'(.+): до ([\d\s]+) ₽ за рабочий день',text(link))
            assert match,'Unrecognized delivery production capacity: '+text(link)
            capacity.append({'name':match[1],'route':link.get('data-route'),'amount':int(re.sub(r'\s','',match[2]))})
        assert len(capacity)==6
        pages['company/delivery/index.html']['production_capacity']=capacity
requirements=[]
for d in load('company/requirements/index.html').xpath('//main//details'):
    route=d.xpath('.//a[@data-route]')[0].get('data-route')
    service=services[route]
    requirements.append({'id':service['id'],'name':service['name'],'category':service['category'],
                         'text':' '.join(text(p) for p in d.xpath('./p[not(a)]'))})
assert len(requirements)==len(catalog)==55
assert {r['id'] for r in requirements}=={s['id'] for s in catalog},'Service requirements coverage changed'
technical_topics=[]
technical_file=source/'technical-data.js'
if technical_file.exists():
    # Parse the customer's JSON literal as content; never execute archive JavaScript.
    prefix='window.TEKST_REQUIREMENTS='
    code=technical_file.read_text(encoding='utf-8')
    assert code.startswith(prefix),'Unexpected technical data format'
    technical,_=json.JSONDecoder().raw_decode(code[len(prefix):])
    assert set(technical['services'])=={s['id'] for s in catalog}
    keys=['send','checks','example','scheme','fixed','note','fileSpec','critical','mistakes','preflight']
    for item in requirements:
        profile=technical['services'][item['id']]
        assert profile['url']==next(s['url'] for s in catalog if s['id']==item['id'])
        assert all(profile.get(key) for key in ['send','checks','example','fileSpec','critical','mistakes','preflight'])
        item['technical']={key:profile[key] for key in keys}
    technical_topics=[{'id':key,'title':value['title'],'text':value['text']} for key,value in technical['topics'].items()]
questions=[]
for h in load('company/vopros-otvet/index.html').xpath('//main//h2'):
    answer=h.getnext()
    assert answer.tag=='p'
    questions.append({'question':text(h),'answer':text(answer),'html':clean(answer)})

articles=[]
for card in load('company/articles/index.html').xpath('//article[@data-article]'):
    link=card.xpath('.//h3/a')[0];route=link.get('data-route')
    article=load(route.strip('/')+'/index.html').xpath('//main/article|//main//article[contains(@class,"reading")]')[0]
    title=text(article.xpath('./h1')[0])
    path=legacy_paths.get(norm(title),'company/article/'+route.strip('/').split('/')[-1]+'/')
    blocks=[];related_services=[]
    for child in article:
        if child.tag=='h2' and text(child)=='Перейти к заказу':break
        if child.tag=='div' and child.xpath('./table'):
            for table in child.xpath('./table'):
                blocks.append({'tag':'table','html':clean(table),'text':text(table)})
            continue
        if child.tag not in ALLOWED or 'article-date' in child.get('class',''):continue
        blocks.append({'tag':child.tag,'html':clean(child),'text':text(child)})
    for link in article.xpath('./div[contains(@class,"service-index")]//a'):
        related_services.append({'route':link.get('data-route'),'label':text(link)})
    dates=article.xpath('./p[contains(@class,"article-date")]')
    category=text(card.xpath('./span')[0])
    if category=='Заказ и макеты':category='Заказ и подготовка макетов'
    record={'id':int(card.get('data-article')),'title':title,'category':category,
                     'path':path,'source_route':route,'teaser':text(card.xpath('./p')[0]),
                     'dates':text(dates[0]) if dates else '', 'blocks':blocks,'services':related_services}
    if not args.seo:
        original=previous_articles[route]
        assert norm(original['title'])==norm(title),'Article title changed; provide its SEO workbook'
        assert original['dates']==record['dates'],'Article dates changed; provide its SEO workbook'
        record['seo']=original['seo']
    articles.append(record)
assert len(articles)==150 and len({a['path'] for a in articles})==150
assert all(a['blocks'] for a in articles)
assert all(a['path'] in {x['path'] for x in articles} for a in legacy)
provenance={'archive':args.archive.name if args.archive else source.name,'imported':args.imported,'company_pages':len(pages),
            'preserved_article_paths':[a['path'] for a in legacy if a['id']<=50]}
if args.archive:provenance['sha256']=hashlib.sha256(args.archive.read_bytes()).hexdigest()
if args.seo:
    from openpyxl import load_workbook
    workbook=load_workbook(args.seo,read_only=True,data_only=True)
    content_rows=list(workbook['05_КОНТЕНТ_55'].iter_rows(values_only=True))
    content_by_id={str(row[0]):dict(zip(content_rows[0],row)) for row in content_rows[1:] if row[0]}
    for requirement in requirements:
        assert norm(requirement['text'])==norm(content_by_id[requirement['id']]['Требования']),requirement['name']+': ZIP/XLSX requirements differ'
    rows=list(workbook['07_СТАТЬИ_150'].iter_rows(values_only=True))
    by_route={row[3]:dict(zip(rows[0],row)) for row in rows[1:] if row[3]}
    assert len(by_route)==len(articles)
    for article in articles:
        row=by_route[article['source_route']]
        assert norm(row['H1 / тема'])==norm(article['title'])
        article['seo']={k:row[k] for k in ['Title','Description','Опубликовано','Обновлено','Услуги ID']}
    provenance['seo_workbook']=args.seo.name
    provenance['seo_sha256']=hashlib.sha256(args.seo.read_bytes()).hexdigest()
    provenance['seo_article_rows_verified']=len(articles)
    provenance['seo_requirement_rows_verified']=len(requirements)
    workbook.close()
if not args.seo:
    for key in ['seo_workbook','seo_sha256','seo_article_rows_verified']:
        if key in previous['source']:provenance[key]=previous['source'][key]
    provenance['seo_article_metadata_retained']=True
if technical_topics:
    provenance['requirements_version']='2026-10-07-tech-deep'
    provenance['technical_profiles_verified']=len(requirements)
    provenance['technical_data_sha256']=hashlib.sha256(technical_file.read_bytes()).hexdigest()
    for name in ['Manrope-OFL.txt','PlayfairDisplay-OFL.txt']:
        destination=ROOT/'fonts'/name
        destination.parent.mkdir(exist_ok=True)
        license_text=(source/'fonts'/name).read_text(encoding='utf-8')
        destination.write_text('\n'.join(line.rstrip() for line in license_text.splitlines())+'\n',encoding='utf-8',newline='\n')
data={'source':provenance,'pages':pages,'requirements':requirements,'technical_topics':technical_topics,'questions':questions,'articles':articles}
(ROOT/'company-content.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(f'Imported {len(pages)} company pages, {len(requirements)} service requirements and {len(articles)} full articles; existing article URLs preserved.')
